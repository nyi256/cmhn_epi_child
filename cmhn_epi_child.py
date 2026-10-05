import io
import re
from datetime import datetime
import numpy as np
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st

# ==========================================
# 1. MASTER IDEAL VARIABLES (Global)
# ==========================================
COLS = [
    "ORGANIZATION",
    "TOWNSHIP_MIMU",
    "TOWNSHIP",
    "CLINIC_NAME",
    "CHANNEL",
    "REPORTING_MONTH",
    "REPORTING_YEAR",
    "SR_NO",
    "FSD",
    "FSD_month",
    "FSD_age_in_months",
    "CHILD_ID",
    "CHILD_NAME",
    "GENDER",
    "DISABILITY_YN",
    "IDP_YN",
    "ADDRESS",
    "DOB",
    "BCG",
    "BCG_REPORTING_PERIOD",
    "OPV1",
    "OPV1_REPORTING_PERIOD",
    "OPV2",
    "OPV2_REPORTING_PERIOD",
    "OPV3",
    "OPV3_REPORTING_PERIOD",
    "Penta1",
    "Penta1_REPORTING_PERIOD",
    "Penta2",
    "Penta2_REPORTING_PERIOD",
    "Penta3",
    "Penta3_REPORTING_PERIOD",
    "IPV",
    "IPV_REPORTING_PERIOD",
    "MMR1",
    "MMR1_REPORTING_PERIOD",
    "MMR2",
    "MMR2_REPORTING_PERIOD",
    "JE1",
    "JE1_REPORTING_PERIOD",
    "JE2",
    "JE2_REPORTING_PERIOD",
    "PCV1",
    "PCV1_REPORTING_PERIOD",
    "PCV2",
    "PCV2_REPORTING_PERIOD",
    "PCV3",
    "PCV3_REPORTING_PERIOD",
    "Rota1",
    "Rota1_REPORTING_PERIOD",
    "Rota2",
    "Rota2_REPORTING_PERIOD",
    "HPV1",
    "HPV1_REPORTING_PERIOD",
    "HPV2",
    "HPV2_REPORTING_PERIOD",
    "REMARK",
    "New_Comer_YN",
    "At Least One Dose Date",
    "At Least One Dose Month",
]

if "historical_children" not in st.session_state:
  st.session_state["historical_children"] = {}

PERIOD_COLS = [
    "BCG_REPORTING_PERIOD",
    "OPV1_REPORTING_PERIOD",
    "OPV2_REPORTING_PERIOD",
    "OPV3_REPORTING_PERIOD",
    "Penta1_REPORTING_PERIOD",
    "Penta2_REPORTING_PERIOD",
    "Penta3_REPORTING_PERIOD",
    "IPV_REPORTING_PERIOD",
    "MMR1_REPORTING_PERIOD",
    "MMR2_REPORTING_PERIOD",
    "JE1_REPORTING_PERIOD",
    "JE2_REPORTING_PERIOD",
    "PCV1_REPORTING_PERIOD",
    "PCV2_REPORTING_PERIOD",
    "PCV3_REPORTING_PERIOD",
    "Rota1_REPORTING_PERIOD",
    "Rota2_REPORTING_PERIOD",
    "HPV1_REPORTING_PERIOD",
    "HPV2_REPORTING_PERIOD",
]


# ==========================================
# 2. UNIVERSAL HELPERS
# ==========================================
def format_date(v):
  if pd.isna(v) or str(v).strip() == "":
    return ""
  val = str(v).strip()
  if val.lower() == "yes":
    return "Other Source"
  d = pd.to_datetime(v, errors="coerce")
  return d.strftime("%d/%m/%Y") if pd.notna(d) else val


def format_period(v):
  if pd.isna(v) or str(v).strip() == "":
    return ""
  d = pd.to_datetime(v, errors="coerce")
  return d.strftime("%b-%y") if pd.notna(d) else str(v).strip()


def clean_year(v):
  if pd.isna(v) or str(v).strip() == "":
    return ""
  val_str = str(v).strip()
  if val_str.endswith(".0"):
    val_str = val_str[:-2]
  return val_str


def finalize_df(df):
  for c in COLS:
    if c not in df.columns:
      df[c] = ""
  return df[COLS].fillna("").replace(["nan", "None", "NaT", "NaN"], "")


def calculate_newcomer_quarter(rec, org_prefix):
  # 🔴 CHANGE THIS NUMBER WHEN THE YEAR CHANGES (e.g., 2027)
  TARGET_YEAR = 2026  
  
  row_dates = []

  for col in PERIOD_COLS:
    val = str(rec.get(col, "")).strip().lower()
    if val and val not in [
        "",
        "nan",
        "none",
        "yes",
        "other source",
        "prevacc",
        "notvacc",
    ]:
      # Parse dates with MFM's formatting approach
      parsed_dt = pd.to_datetime(val, format="%b-%y", errors="coerce")
      if pd.isna(parsed_dt):
        parsed_dt = pd.to_datetime(val, errors="coerce", dayfirst=True)

      # Only accept dates that match the TARGET_YEAR exactly
      if pd.notna(parsed_dt) and parsed_dt.year == TARGET_YEAR:
        row_dates.append(parsed_dt)

  # Assign Quarter based on the earliest valid date found
  if row_dates:
    earliest_vax = min(row_dates)
    calc_q = (earliest_vax.month - 1) // 3 + 1
    assigned_q = f"Q{calc_q} {TARGET_YEAR}"
  else:
    assigned_q = ""

  # ID standardization
  cid_str = str(rec.get("CHILD_ID", "")).strip().lower()
  if cid_str.endswith(".0"):
    cid_str = cid_str[:-2]
  cname = str(rec.get("CHILD_NAME", "")).strip().lower().replace(" ", "")

  # Stricter unique key combining ID and Name
  if cid_str in ["", "nan", "none", "///", "//", "-", "_"]:
    unique_key = f"{org_prefix}_missingid_{cname}_{rec.get('SR_NO','')}"
  else:
    unique_key = f"{org_prefix}_{cid_str}_{cname}"

  # Historical Validation Return Logic
  if unique_key in st.session_state["historical_children"]:
    earliest_assigned = st.session_state["historical_children"][unique_key]
    return earliest_assigned if assigned_q == earliest_assigned else ""
  else:
    if assigned_q:
      st.session_state["historical_children"][unique_key] = assigned_q
    return assigned_q


def calculate_months_diff(dob, target_dt):
  if pd.isna(dob) or pd.isna(target_dt):
    return np.nan
  months = (target_dt.year - dob.year) * 12 + (target_dt.month - dob.month)
  if target_dt.day < dob.day:
    months -= 1
  return months


# ==========================================
# 3. MULTI-QUARTER LOOKUP ENGINE
# ==========================================
def extract_multi_quarter_lookup(previous_files):
  if not previous_files:
    return {}

  master_lookup = {}
  for prev_file in previous_files:
    try:
      xls = pd.ExcelFile(prev_file)
      sheet = next(
          (
              s
              for s in [
                  "EPIC_Child",
                  "EPI_Data Entry_Child_Share",
                  "EPI_child",
              ]
              if s in xls.sheet_names
          ),
          xls.sheet_names[0],
      )
      df_raw = pd.read_excel(prev_file, sheet_name=sheet, header=None)

      header_idx = next(
          (
              i
              for i in range(min(20, len(df_raw)))
              if any(
                  "date" in str(x).lower() and "vaccination" in str(x).lower()
                  for x in df_raw.iloc[i].values
                  if pd.notna(x)
              )
          ),
          1,
      )
      h1 = df_raw.iloc[header_idx].fillna("").astype(str).tolist()
      h2 = df_raw.iloc[header_idx + 1].fillna("").astype(str).tolist()

      combined_headers = []
      curr_main = ""
      for main, sub in zip(h1, h2):
        if main.strip():
          curr_main = main
        combined_headers.append(
            f"{curr_main} {sub}"
            .strip()
            .replace("\n", " ")
            .replace("\r", "")
            .lower()
            .replace(" ", "")
        )

      df = df_raw.iloc[header_idx + 2 :].reset_index(drop=True)
      idx_id = next(
          (
              idx
              for idx, h in enumerate(combined_headers)
              if "registerno" in h or "childid" in h
          ),
          None,
      )

      round_col_map = {
          int(re.search(r"round(\d+)", h).group(1)): idx
          for idx, h in enumerate(combined_headers)
          if "dateofvaccination" in h and re.search(r"round(\d+)", h)
      }
      if (
          2 in round_col_map
          and 3 not in round_col_map
          and round_col_map[2] + 1 < len(combined_headers)
      ):
        round_col_map[3] = round_col_map[2] + 1

      if idx_id is not None:
        for _, row in df.iterrows():
          cid = clean_year(row.iloc[idx_id]).strip()
          if not cid:
            continue
          if cid not in master_lookup:
            master_lookup[cid] = {}

          for r_num, c_idx in round_col_map.items():
            dt_val = row.iloc[c_idx]
            parsed_dt = pd.to_datetime(dt_val, errors="coerce")
            if pd.notna(parsed_dt):
              master_lookup[cid][r_num] = parsed_dt
    except Exception as e:
      st.warning(f"Error parsing previous file {prev_file.name}: {e}")

  return master_lookup


def process_cmhn(uploaded_file, sheet_name=None, history_lookup=None):
  xls = pd.ExcelFile(uploaded_file)
  actual_sheet = (
      sheet_name
      if sheet_name in xls.sheet_names
      else next(
          (
              s
              for s in [
                  "EPIC_Child",
                  "EPI_Data Entry_Child_Share",
                  "EPI_child",
              ]
              if s in xls.sheet_names
          ),
          xls.sheet_names[0],
      )
  )
  df_raw = pd.read_excel(uploaded_file, sheet_name=actual_sheet, header=None)

  header_idx = next(
      (
          i
          for i in range(min(20, len(df_raw)))
          if any(
              "date" in str(x).lower() and "vaccination" in str(x).lower()
              for x in df_raw.iloc[i].values
              if pd.notna(x)
          )
      ),
      1,
  )
  h1 = df_raw.iloc[header_idx].fillna("").astype(str).tolist()
  h2 = df_raw.iloc[header_idx + 1].fillna("").astype(str).tolist()

  combined_headers = []
  curr_main = ""
  for main, sub in zip(h1, h2):
    if main.strip():
      curr_main = main
    combined_headers.append(
        f"{curr_main} {sub}"
        .strip()
        .replace("\n", " ")
        .replace("\r", "")
        .lower()
        .replace(" ", "")
    )

  df = df_raw.iloc[header_idx + 2 :].reset_index(drop=True)
  round_col_map = {
      int(re.search(r"round(\d+)", h).group(1)): idx
      for idx, h in enumerate(combined_headers)
      if "dateofvaccination" in h and re.search(r"round(\d+)", h)
  }
  if (
      2 in round_col_map
      and 3 not in round_col_map
      and round_col_map[2] + 1 < len(combined_headers)
  ):
    round_col_map[3] = round_col_map[2] + 1

  def get_col_idx(keywords):
    return next(
        (
            idx
            for idx, h in enumerate(combined_headers)
            if all(k.lower() in h for k in keywords)
        ),
        None,
    )

  idx_tsp = get_col_idx(["township"])
  idx_clinic = get_col_idx(["session"]) or get_col_idx(["teamname"])
  idx_id = get_col_idx(["registerno"]) or get_col_idx(["childid"])
  idx_gender, idx_address = get_col_idx(["gender"]), get_col_idx(["address"])
  idx_dis, idx_idp, idx_dob = (
      get_col_idx(["disability"]),
      get_col_idx(["idp"]),
      get_col_idx(["birth"]) or get_col_idx(["dob"]),
  )

  vax_map = {
      "BCG": get_col_idx(["bcg"]),
      "Penta1": get_col_idx(["penta", "1st"]),
      "Penta2": get_col_idx(["penta", "2nd"]),
      "Penta3": get_col_idx(["penta", "3rd"]),
      "OPV1": get_col_idx(["opv", "1st"]),
      "OPV2": get_col_idx(["opv", "2nd"]),
      "OPV3": get_col_idx(["opv", "3rd"]),
      "IPV": get_col_idx(["ipv"]),
      "MMR1": get_col_idx(["mmr", "1st"]),
      "MMR2": get_col_idx(["mmr", "2nd"]),
      "JE1": get_col_idx(["je", "1st"]),
      "JE2": get_col_idx(["je", "2nd"]),
      "PCV1": get_col_idx(["pcv", "1st"]),
      "PCV2": get_col_idx(["pcv", "2nd"]),
      "PCV3": get_col_idx(["pcv", "3rd"]),
      "Rota1": get_col_idx(["rota", "1st"]),
      "Rota2": get_col_idx(["rota", "2nd"]),
      "HPV1": get_col_idx(["hpv", "1st"]),
      "HPV2": get_col_idx(["hpv", "2nd"]),
  }

  records = []
  for _, row in df.iterrows():
    parsed_round_dates = {}
    for r_num, c_idx in round_col_map.items():
      raw_val = str(row.iloc[c_idx]).strip()
      if raw_val.lower() not in ["nan", "", "none", "nat"]:
        match = re.search(r"(\d{4}-\d{2}-\d{2})", raw_val)
        parsed_round_dates[r_num] = pd.to_datetime(
            match.group(1) if match else raw_val, errors="coerce"
        )

    id_val = clean_year(row.iloc[idx_id]) if idx_id is not None else ""

    if history_lookup and id_val in history_lookup:
      for r_k, r_v in history_lookup[id_val].items():
        if r_k not in parsed_round_dates:
          parsed_round_dates[r_k] = r_v

    master_dt = (
        parsed_round_dates.get(2)
        or parsed_round_dates.get(3)
        or parsed_round_dates.get(1)
        or next((d for d in parsed_round_dates.values() if d is not None), None)
    )
    if id_val.lower() in ["nan", "", "none"] and pd.isna(master_dt):
      continue

    rec = {
        "ORGANIZATION": "CMHN",
        "TOWNSHIP_MIMU": (
            str(row.iloc[idx_tsp]).strip() if idx_tsp is not None else ""
        ),
        "TOWNSHIP": (
            str(row.iloc[idx_tsp]).strip() if idx_tsp is not None else ""
        ),
        "CLINIC_NAME": (
            str(row.iloc[idx_clinic]).strip() if idx_clinic is not None else ""
        ),
        "REPORTING_MONTH": master_dt.strftime("%b") if pd.notna(master_dt) else "",
        "REPORTING_YEAR": (
            clean_year(master_dt.year) if pd.notna(master_dt) else ""
        ),
        "CHILD_ID": id_val,
        "CHILD_NAME": "",
        "GENDER": (
            str(row.iloc[idx_gender]).strip() if idx_gender is not None else ""
        ),
        "DISABILITY_YN": (
            str(row.iloc[idx_dis]).strip() if idx_dis is not None else "N"
        ),
        "IDP_YN": str(row.iloc[idx_idp]).strip() if idx_idp is not None else "N",
        "ADDRESS": (
            str(row.iloc[idx_address]).strip()
            if idx_address is not None
            else ""
        ),
        "DOB": format_date(row.iloc[idx_dob]) if idx_dob is not None else "",
    }

    for vax_name, v_idx in vax_map.items():
      if v_idx is not None:
        val = str(row.iloc[v_idx]).strip().lower().replace(" ", "")
        if "prev" in val:
          rec[vax_name] = "Other Source"
        elif "round" in val:
          vax_match = re.search(r"round(\d+)", val)
          if vax_match:
            vax_round_num = int(vax_match.group(1))
            target_dt = parsed_round_dates.get(vax_round_num)
            if target_dt and pd.notna(target_dt):
              rec[vax_name] = target_dt.strftime("%d/%m/%Y")
              rec[f"{vax_name}_REPORTING_PERIOD"] = target_dt.strftime("%b-%y")
            else:
              rec[vax_name] = f"Round {vax_round_num}"
              if vax_round_num == 1:
                rec[f"{vax_name}_REPORTING_PERIOD"] = "Mar-26"
              elif vax_round_num == 2:
                rec[f"{vax_name}_REPORTING_PERIOD"] = "Apr-26"
              elif vax_round_num == 3:
                rec[f"{vax_name}_REPORTING_PERIOD"] = "May-26"
              else:
                rec[f"{vax_name}_REPORTING_PERIOD"] = ""

    rec["New_Comer_YN"] = calculate_newcomer_quarter(rec, "CMHN")
    records.append(rec)

  return finalize_df(pd.DataFrame(records))


# ==========================================
# 4. EXCEL STYLING & EXPORT
# ==========================================
def export_excel(df, partner_name):
  output = io.BytesIO()
  with pd.ExcelWriter(output, engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name="Register")
    ws = writer.sheets["Register"]
    header_fill = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid"
    )
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center")
    for cell in ws[1]:
      cell.fill = header_fill
      cell.font = header_font
      cell.alignment = header_alignment
    ws.freeze_panes = "A2"
    for col in ws.columns:
      ws.column_dimensions[get_column_letter(col[0].column)].width = 15
  return output.getvalue()


# ==========================================
# 5. VALIDATION ENGINE (With Mutually Exclusive Full Dose & FSD)
# ==========================================
def process_data_validation(df):
  if df is None or df.empty:
    return None, None
  vdf = df.copy()

  vax_list = [
      "BCG",
      "OPV1",
      "OPV2",
      "OPV3",
      "Penta1",
      "Penta2",
      "Penta3",
      "IPV",
      "MMR1",
      "MMR2",
      "JE1",
      "JE2",
      "PCV1",
      "PCV2",
      "PCV3",
      "Rota1",
      "Rota2",
      "HPV1",
      "HPV2",
  ]

  def parse_single_date(val):
    if pd.isna(val) or str(val).strip().lower() in [
        "",
        "nan",
        "other source",
        "prevacc",
        "notvacc",
        "round 1",
        "round 2",
        "round 3",
        "round 4",
        "round 5",
    ]:
      return pd.NaT
    val_str = str(val).strip()
    try:
      if "-" in val_str and val_str[:4].isdigit():
        return pd.to_datetime(val_str, errors="coerce")
      return pd.to_datetime(val_str, dayfirst=True, errors="coerce")
    except Exception:
      return pd.NaT

  if "DOB" in vdf.columns:
    vdf["DOB_dt"] = vdf["DOB"].apply(parse_single_date)
  else:
    vdf["DOB_dt"] = pd.NaT

  for v in vax_list:
    if v in vdf.columns:
      vdf[f"{v}_dt"] = vdf[v].apply(parse_single_date)

      def get_precise_age(row, v_dt_col, vax_name):
        if pd.isna(row["DOB_dt"]) or pd.isna(row[v_dt_col]):
          return np.nan
        dob = row["DOB_dt"]
        vax_date = row[v_dt_col]

        months_diff = (vax_date.year - dob.year) * 12 + (
            vax_date.month - dob.month
        )
        if vax_date.day < dob.day:
          months_diff -= 1

        if vax_name in ["Penta1", "Penta2", "Penta3", "OPV1", "OPV2", "OPV3"]:
          try:
            milestone_date = dob + pd.DateOffset(months=months_diff)
          except ValueError:
            milestone_date = dob + pd.offsets.MonthEnd(months_diff)
          remaining_days = (vax_date - milestone_date).days
          return round(months_diff + (remaining_days / 100), 2)
        else:
          return int(months_diff)

      vdf[f"{v}_Age"] = vdf.apply(
          lambda r: get_precise_age(r, f"{v}_dt", v), axis=1
      )

  # --- FSD (First Service Date) LOGIC ---
  dt_cols = [f"{v}_dt" for v in vax_list if f"{v}_dt" in vdf.columns]
  vdf["FSD_dt"] = vdf[dt_cols].min(axis=1)

  vdf["FSD"] = vdf["FSD_dt"].apply(
      lambda d: d.strftime("%d/%m/%Y") if pd.notna(d) else ""
  )
  vdf["FSD_month"] = vdf["FSD_dt"].apply(
      lambda d: d.strftime("%b-%y") if pd.notna(d) else ""
  )
  vdf["FSD_age_in_months"] = vdf.apply(
      lambda r: str(int(calculate_months_diff(r["DOB_dt"], r["FSD_dt"])))
      if pd.notna(r["DOB_dt"]) and pd.notna(r["FSD_dt"]) else "",
      axis=1
  )

  def get_valid_last_visit(row):
    valid_dates = []
    for v in vax_list:
      dt_col = f"{v}_dt"
      rp_col = f"{v}_REPORTING_PERIOD"
      if (
          dt_col in row
          and pd.notna(row[dt_col])
          and rp_col in row
          and pd.notna(row[rp_col])
          and str(row[rp_col]).strip() != ""
      ):
        valid_dates.append(row[dt_col])
    return max(valid_dates) if valid_dates else pd.NaT

  vdf["Vaccination Date (last visit)_dt"] = vdf.apply(
      get_valid_last_visit, axis=1
  )
  vdf["Age in Months (Last visit)"] = vdf.apply(
      lambda r: calculate_months_diff(
          r["DOB_dt"], r["Vaccination Date (last visit)_dt"]
      ),
      axis=1,
  )
  vdf["Vaccination Date (last visit)"] = vdf[
      "Vaccination Date (last visit)_dt"
  ].apply(lambda d: d.strftime("%d/%m/%Y") if pd.notna(d) else "")

  period_cols = [c for c in vdf.columns if "_REPORTING_PERIOD" in c]
  vdf["At Least One Dose"] = vdf.apply(
      lambda row: (
          "Yes"
          if "".join([
              str(row[c]) for c in period_cols if pd.notna(row[c])
          ]).strip()
          != ""
          else "No"
      ),
      axis=1,
  )

  QUARTER_CONFIG = {
      "Q1 2026": {
          "start": pd.Timestamp("2026-01-01"),
          "end": pd.Timestamp("2026-03-31"),
          "months": ["Jan-26", "Feb-26", "Mar-26"],
      },
      "Q2 2026": {
          "start": pd.Timestamp("2026-04-01"),
          "end": pd.Timestamp("2026-06-30"),
          "months": ["Apr-26", "May-26", "Jun-26"],
      },
      "Q3 2026": {
          "start": pd.Timestamp("2026-07-01"),
          "end": pd.Timestamp("2026-09-30"),
          "months": ["Jul-26", "Aug-26", "Sep-26"],
      },
      "Q4 2026": {
          "start": pd.Timestamp("2026-10-01"),
          "end": pd.Timestamp("2026-12-31"),
          "months": ["Oct-26", "Nov-26", "Dec-26"],
      },
  }

  def evaluate_cohort_at_least_one(row):
    default_res = pd.Series(
        ["", ""], index=["At Least One Dose Date", "At Least One Dose Month"]
    )
    raw_cohort = row.get("New_Comer_YN", None)

    if (
        pd.isna(raw_cohort)
        or str(raw_cohort).strip() == ""
        or str(raw_cohort).strip().lower() in ["nan", "none"]
    ):
      return default_res

    cohort = str(raw_cohort).strip()
    if cohort not in QUARTER_CONFIG:
      return default_res

    q_cfg = QUARTER_CONFIG[cohort]
    q_start, q_end = q_cfg["start"], q_cfg["end"]
    q_target_months = q_cfg["months"]

    cohort_dates = []
    for v in vax_list:
      rp_val = str(row.get(f"{v}_REPORTING_PERIOD", "")).strip()
      dt_val = row.get(f"{v}_dt", pd.NaT)
      if pd.notna(dt_val):
        if (q_start <= dt_val <= q_end) or (rp_val in q_target_months):
          cohort_dates.append(dt_val)

    if not cohort_dates:
      return default_res

    latest_cohort_date = max(cohort_dates)
    at_least_one_date_str = latest_cohort_date.strftime("%d/%m/%Y")

    dob_dt = row.get("DOB_dt", pd.NaT)
    m_diff = calculate_months_diff(dob_dt, latest_cohort_date)
    at_least_one_age_m = int(m_diff) if pd.notna(m_diff) else ""

    return pd.Series(
        [at_least_one_date_str, at_least_one_age_m],
        index=["At Least One Dose Date", "At Least One Dose Month"],
    )

  if not vdf.empty:
    vdf[["At Least One Dose Date", "At Least One Dose Month"]] = vdf.apply(
        evaluate_cohort_at_least_one, axis=1
    )

  # --- STRICTLY MUTUALLY EXCLUSIVE FULL DOSE LOGIC ---
  def check_full_dose_u1(row):
    core_vax_cols = [
        "BCG", "OPV1", "OPV2", "OPV3", "Penta1", "Penta2", "Penta3", "MMR1",
    ]
    fsd_age = row.get("FSD_age_in_months", "")
    if pd.isna(fsd_age) or str(fsd_age).strip() == "":
      return ""
    if float(fsd_age) >= 12:
      return ""  
    if str(row.get("At Least One Dose", "")).strip() != "Yes":
      return ""
      
    all_present = all(
        str(row.get(c, "")).strip() != "" for c in core_vax_cols if c in row.index
    )
    mmr1_age_val = row.get("MMR1_Age", None)
    mmr1_str = str(row.get("MMR1", "")).strip().lower()
    mmr1_valid = False
    
    if mmr1_str == "other source":
      last_visit_age = row.get("Age in Months (Last visit)", np.nan)
      mmr1_valid = pd.notna(last_visit_age) and float(last_visit_age) < 12
    elif pd.notna(mmr1_age_val) and str(mmr1_age_val).strip() != "":
      try:
        mmr1_valid = 9 <= float(mmr1_age_val) < 12
      except (ValueError, TypeError):
        pass

    if all_present and mmr1_valid:
      return "Full Dose"
    return ""

  def check_full_dose_1to5(row):
    fsd_age = row.get("FSD_age_in_months", "")
    if pd.isna(fsd_age) or str(fsd_age).strip() == "":
      return ""

    if float(fsd_age) < 12:
      mmr1_age_val = row.get("MMR1_Age", None)
      mmr1_str = str(row.get("MMR1", "")).strip().lower()
      last_visit = row.get("Age in Months (Last visit)", np.nan)
      is_delayed = False
      if mmr1_str == "other source":
        is_delayed = pd.notna(last_visit) and float(last_visit) >= 12
      elif pd.notna(mmr1_age_val):
        try:
          is_delayed = float(mmr1_age_val) >= 12
        except (ValueError, TypeError):
          pass
      if not is_delayed:
        return ""

    core_vax_cols = ["OPV1", "OPV2", "OPV3", "Penta1", "Penta2", "Penta3", "MMR1"]
    if str(row.get("At Least One Dose", "")).strip() != "Yes":
      return ""
      
    all_present = all(
        str(row.get(c, "")).strip() != "" for c in core_vax_cols if c in row.index
    )
    mmr1_age_val = row.get("MMR1_Age", None)
    mmr1_str = str(row.get("MMR1", "")).strip().lower()
    mmr1_valid = False
    
    if mmr1_str == "other source":
      mmr1_valid = True
    elif pd.notna(mmr1_age_val) and str(mmr1_age_val).strip() != "":
      try:
        mmr1_valid = float(mmr1_age_val) >= 9
      except (ValueError, TypeError):
        pass

    if all_present and mmr1_valid:
      return "Full Dose"
    return ""

  vdf["REACH Full Dose <1"] = vdf.apply(check_full_dose_u1, axis=1)
  vdf["REACH Full Dose 1 To 5"] = vdf.apply(check_full_dose_1to5, axis=1)
  vdf["Full Dose_U5"] = np.where(
      (vdf["REACH Full Dose <1"] == "Full Dose")
      | (vdf["REACH Full Dose 1 To 5"] == "Full Dose"),
      "Full Dose",
      "",
  )

  core_rp_list = [
      "OPV1_REPORTING_PERIOD", "OPV2_REPORTING_PERIOD", "OPV3_REPORTING_PERIOD",
      "Penta1_REPORTING_PERIOD", "Penta2_REPORTING_PERIOD", "Penta3_REPORTING_PERIOD",
      "MMR1_REPORTING_PERIOD",
  ]
  temp_date_df = pd.DataFrame(index=vdf.index)
  for rp_col in core_rp_list:
    if rp_col in vdf.columns:
      temp_date_df[rp_col] = pd.to_datetime(
          vdf[rp_col], format="%b-%y", errors="coerce"
      )

  latest_calendar_dates = (
      temp_date_df.max(axis=1)
      if not temp_date_df.empty
      else pd.Series(pd.NaT, index=vdf.index)
  )
  vdf["Full Dose Reporting Month"] = latest_calendar_dates.dt.strftime(
      "%b-%y"
  ).fillna("")
  vdf.loc[vdf["Full Dose_U5"] != "Full Dose", "Full Dose Reporting Month"] = ""

  errors = []
  for v in vax_list:
    if f"{v}_dt" in vdf.columns:
      mask = (vdf[f"{v}_dt"] < vdf["DOB_dt"]) & vdf[f"{v}_dt"].notna()
      if mask.any():
        err_df = vdf[mask].copy()
        err_df["Error_Details"] = f"{v} before DOB"
        errors.append(err_df)

  vdf["DOB"] = vdf["DOB_dt"].dt.strftime("%d/%m/%Y")
  for v in vax_list:
    if v in vdf.columns and f"{v}_dt" in vdf.columns:
      vdf[v] = vdf.apply(
          lambda r: (
              r[f"{v}_dt"].strftime("%d/%m/%Y")
              if pd.notnull(r[f"{v}_dt"])
              else r[v]
          ),
          axis=1,
      )

  final_vdf = vdf.drop(columns=[c for c in vdf.columns if c.endswith("_dt")])
  mistakes_df = pd.concat(errors).drop_duplicates() if errors else pd.DataFrame()

  return final_vdf, mistakes_df


# ==========================================
# 6. MAIN NAVIGATION
# ==========================================
st.set_page_config(page_title="EPI Multi-Year Master", layout="wide")
st.sidebar.title("💉 EPI Navigation")
page = st.sidebar.radio(
    "Go to:",
    [
        "Data Processing",
        "Data Validation",
        "EPIC Indicator",
    ],
)

if page == "Data Processing":
  st.title("💉 EPI Partner - CMHN Multi-Quarter Processing")

  col1, col2 = st.columns(2)
  with col1:
    file_cmhn = st.file_uploader(
        "Upload Current Quarter File (e.g. Q2, Q3, Q4 2026, Q1 2027...)",
        type=["xlsx"],
        key="cmhn_up",
    )
  with col2:
    files_baseline = st.file_uploader(
        "Upload All Previous Files (Q1, Q2, Q3, etc.) for Cumulative Baseline",
        type=["xlsx"],
        accept_multiple_files=True,
        key="cmhn_baseline_up",
    )

  if file_cmhn:
    xl_cmhn = pd.ExcelFile(file_cmhn)
    sheet_cmhn = st.selectbox(
        "Select Sheet to Process", xl_cmhn.sheet_names, key="cmhn_sheet_sel"
    )

    if st.button("🚀 Process & Link All Quarters", key="cmhn_btn"):
      with st.spinner("Extracting historical lookup and linking timelines..."):
        history_lookup = (
            extract_multi_quarter_lookup(files_baseline)
            if files_baseline
            else {}
        )
        df_cmhn = process_cmhn(
            file_cmhn, sheet_cmhn, history_lookup=history_lookup
        )
        st.session_state["cmhn_processed"] = df_cmhn
        final_v, _ = process_data_validation(df_cmhn)
        st.session_state["cmhn_processed_v_final"] = final_v
        st.success(
            "✅ Successfully processed and synchronized across all quarter"
            " baselines!"
        )

  st.divider()
  if "cmhn_processed" in st.session_state:
    st.subheader("📊 Processed Preview")
    combined = st.session_state["cmhn_processed"]
    st.dataframe(combined.head(50), use_container_width=True)
    st.download_button(
        "🚀 Download Master Report",
        export_excel(combined, "CMHN"),
        "CMHN_Master_Report.xlsx",
    )

elif page == "Data Validation":
  st.title("🛡️ Data Validation (Master Log)")
  partners = [("CMHN", "cmhn_processed")]

  st.header("📊 Data Preview")
  for name, key in partners:
    st.subheader(f"🏢 {name}")
    if key in st.session_state:
      final_v, _ = process_data_validation(st.session_state[key])
      st.session_state[f"{key}_v_final"] = final_v

      if final_v is not None and not final_v.empty:
        st.dataframe(final_v.head(20), use_container_width=True)
        st.download_button(
            f"📥 Download {name} Final",
            export_excel(final_v, name),
            f"{name}_Validated.xlsx",
            key=f"vax_dl_{name}",
        )
      else:
        st.warning("⚠️ No validated data available to display.")
    else:
      st.warning(f"No {name} data processed yet.")

  st.divider()

  st.header("🛡️ Partner Validation & Master Log Generation")

  vax_cols_master = [
      "BCG",
      "OPV1",
      "OPV2",
      "OPV3",
      "Penta1",
      "Penta2",
      "Penta3",
      "IPV",
      "MMR1",
      "MMR2",
      "JE1",
      "JE2",
      "PCV1",
      "PCV2",
      "PCV3",
      "Rota1",
      "Rota2",
      "HPV1",
      "HPV2",
  ]

  def generate_composite_key(row):
    cid = str(row.get("CHILD_ID", "")).strip().lower()
    if cid.endswith(".0"):
      cid = cid[:-2]
    if cid in ["nan", "none", "", "nat", "///", "-", "_"]:
      cid = (
          str(row.get("CHILD_NAME", ""))
          .strip()
          .lower()
          .replace(" ", "")
      )
    dob = str(row.get("DOB", "")).strip().lower()
    gender = str(row.get("GENDER", "")).strip().lower()
    err = str(row.get("Error_Details", "")).strip().lower()

    target_date = ""
    for vax in vax_cols_master:
      if vax.lower() in err:
        target_date = str(row.get(vax, "")).strip().lower()
        break

    return f"{cid}_{dob}_{gender}_{err}_{target_date}"

  for name, key in partners:
    st.subheader(f"🏢 {name} Rulebook & Master Log")

    col_upload_1, col_upload_2 = st.columns(2)
    with col_upload_1:
      partner_rule_file = st.file_uploader(
          f"📂 Upload {name} Rulebook", type=["xlsx"], key=f"rule_{name}"
      )
    with col_upload_2:
      historical_master_log = st.file_uploader(
          f"📂 Upload Previous Master Log (Optional for First Run)",
          type=["xlsx"],
          key=f"hist_{name}",
      )

    final_key = f"{key}_v_final"

    if final_key in st.session_state and st.session_state[final_key] is not None:
      v_data = st.session_state[final_key]
      rule_errors_list = []

      if partner_rule_file:
        try:
          rules_df = pd.read_excel(partner_rule_file)

          for _, rule in rules_df.iterrows():
            cat = str(rule.get("Category", "")).strip()
            var = str(rule.get("Variable", "")).strip()
            val = str(rule.get("Rule_Value", "")).strip()

            if cat == "Missing":
              if var in v_data.columns:
                mask = v_data[var].isna() | (
                    v_data[var].astype(str).str.strip() == ""
                )
                bad = v_data[mask].copy()
                if not bad.empty:
                  bad["Error_Details"] = f"Missing {var}"
                  rule_errors_list.append(bad)

            elif cat == "Unique":
              if var in v_data.columns:
                dups = v_data[
                    v_data.duplicated(subset=[var], keep=False)
                    & v_data[var].notna()
                    & (v_data[var].astype(str).str.strip() != "")
                ].copy()
                if not dups.empty:
                  dups["Error_Details"] = f"Duplicate {var}"
                  rule_errors_list.append(dups)

            elif cat == "Value":
              if var in v_data.columns:
                allowed = [v.strip() for v in val.split(",")]
                mask = (
                    ~v_data[var].astype(str).isin(allowed)
                    & v_data[var].notna()
                    & (v_data[var].astype(str).str.strip() != "")
                )
                bad = v_data[mask].copy()
                if not bad.empty:
                  bad["Error_Details"] = f"Invalid {var}: {val}"
                  rule_errors_list.append(bad)

            elif cat == "Age":
              a_col = var if var.endswith("_Age") else f"{var}_Age"
              if a_col in v_data.columns:

                def string_to_days(age_val):
                  if pd.isna(age_val) or str(age_val).strip() in [
                      "",
                      "nan",
                      "None",
                  ]:
                    return None
                  s_val = str(age_val).strip().lower()
                  try:
                    f_val = float(s_val)
                    m = int(f_val)
                    d = round((f_val - m) * 100)
                    return (m * 30.44) + d
                  except ValueError:
                    pass
                  match = re.match(r"(\d+)m\s*(\d+)d", s_val)
                  if match:
                    return (int(match.group(1)) * 30.44) + int(
                        match.group(2)
                    )
                  return None

                try:
                  low_str, high_str = val.split("-")
                  min_limit_val = float(low_str.strip())
                  max_limit_val = float(high_str.strip())
                  min_days = string_to_days(low_str.strip())
                  max_days = string_to_days(high_str.strip())
                except Exception:
                  continue

                v_col_check = var.replace("_Age", "")
                if v_col_check in v_data.columns:
                  vaccine_was_given = v_data[v_col_check].notna() & (
                      v_data[v_col_check].astype(str).str.strip() != ""
                  )
                  active_records = v_data[vaccine_was_given].copy()
                else:
                  active_records = v_data.copy()

                if not active_records.empty:
                  active_records["_temp_days"] = active_records[a_col].apply(
                      string_to_days
                  )
                  violators = active_records[
                      active_records["_temp_days"].notna()
                  ].copy()

                  if not violators.empty:
                    if min_days is not None:
                      bad_below = violators[
                          violators["_temp_days"] < min_days
                      ].copy()
                      if not bad_below.empty:
                        clean_min = (
                            int(min_limit_val)
                            if min_limit_val.is_integer()
                            else min_limit_val
                        )
                        bad_below["Error_Details"] = (
                            f"{var} Age below {clean_min} months"
                        )
                        rule_errors_list.append(
                            bad_below.drop(
                                columns=["_temp_days"], errors="ignore"
                            )
                        )

                    if max_days is not None:
                      bad_over = violators[
                          violators["_temp_days"] > max_days
                      ].copy()
                      if not bad_over.empty:
                        clean_max = (
                            int(max_limit_val)
                            if max_limit_val.is_integer()
                            else max_limit_val
                        )
                        bad_over["Error_Details"] = (
                            f"{var} Age Over {clean_max} months"
                        )
                        rule_errors_list.append(
                            bad_over.drop(
                                columns=["_temp_days"], errors="ignore"
                            )
                        )

            elif cat == "Logic":
              today_ts = pd.Timestamp.now()
              pure_vax_cols = [
                  v for v in vax_cols_master if v in v_data.columns
              ]

              if var in ["All Vaccines", "Future Check"]:
                check_cols = list(pure_vax_cols)
                if "DOB" in v_data.columns:
                  check_cols.append("DOB")

                for col in check_cols:
                  dates = pd.to_datetime(
                      v_data[col], errors="coerce", dayfirst=True
                  )
                  bad = v_data[dates > today_ts].copy()
                  if not bad.empty:
                    bad["Error_Details"] = f"Future Date in {col}"
                    rule_errors_list.append(bad)

              elif var == "DOB_Sequence":
                if "DOB" in v_data.columns:
                  dob_dates = pd.to_datetime(
                      v_data["DOB"], errors="coerce", dayfirst=True
                  )
                  for col in pure_vax_cols:
                    vax_dates = pd.to_datetime(
                        v_data[col], errors="coerce", dayfirst=True
                    )
                    bad = v_data[
                        (vax_dates < dob_dates)
                        & vax_dates.notna()
                        & dob_dates.notna()
                    ].copy()
                    if not bad.empty:
                      bad["Error_Details"] = f"{col} given before birth"
                      rule_errors_list.append(bad)

            elif cat == "Sequence":
              prev_col = val
              if var in v_data.columns and prev_col in v_data.columns:
                current_dates = pd.to_datetime(
                    v_data[var], errors="coerce", dayfirst=True
                )
                prev_dates = pd.to_datetime(
                    v_data[prev_col], errors="coerce", dayfirst=True
                )
                bad = v_data[
                    (current_dates < prev_dates)
                    & current_dates.notna()
                    & prev_dates.notna()
                ].copy()
                if not bad.empty:
                  bad["Error_Details"] = (
                      f"Sequence Error: {var} is before {prev_col}"
                  )
                  rule_errors_list.append(bad)

        except Exception as e:
          st.error(f"Rulebook processing error: {e}")

      if rule_errors_list:
        df_curr = pd.concat(rule_errors_list, ignore_index=True).drop_duplicates()
        df_curr["match_key"] = df_curr.apply(generate_composite_key, axis=1)

        if historical_master_log:
          try:
            df_master = pd.read_excel(historical_master_log)
            if "match_key" not in df_master.columns:
              df_master["match_key"] = df_master.apply(
                  generate_composite_key, axis=1
              )

            curr_keys = set(df_curr["match_key"])
            master_keys = set(df_master["match_key"])

            new_keys = curr_keys - master_keys
            persistent_keys = curr_keys & master_keys
            resolved_keys = master_keys - curr_keys

            df_new = df_curr[df_curr["match_key"].isin(new_keys)].copy()
            df_new["Partner_Response"] = ""
            df_new["Case_Status"] = "Open"
            df_new["Quarter"] = ""
            df_new["System_Flag"] = "Active"

            df_persistent = df_curr[
                df_curr["match_key"].isin(persistent_keys)
            ].copy()
            master_tracking = df_master.set_index("match_key")[
                ["Partner_Response", "Case_Status", "Quarter"]
            ]
            df_persistent = df_persistent.join(
                master_tracking, on="match_key"
            )

            df_persistent["Partner_Response"] = (
                df_persistent["Partner_Response"]
                .fillna("")
                .astype(str)
                .replace("nan", "")
            )
            df_persistent["Case_Status"] = (
                df_persistent["Case_Status"]
                .fillna("Open")
                .astype(str)
                .replace("nan", "Open")
            )
            df_persistent["Quarter"] = (
                df_persistent["Quarter"]
                .fillna("")
                .astype(str)
                .replace("nan", "")
            )
            df_persistent["System_Flag"] = "Active"

            df_resolved = df_master[
                df_master["match_key"].isin(resolved_keys)
            ].copy()
            df_resolved["System_Flag"] = "Fixed/Resolved"

            final_master_log = pd.concat(
                [df_persistent, df_new, df_resolved], ignore_index=True
            )

          except Exception as e:
            st.error(f"Error merging with Master Log: {e}")
            final_master_log = pd.DataFrame()
        else:
          final_master_log = df_curr.copy()
          final_master_log["Partner_Response"] = ""
          final_master_log["Case_Status"] = "Open"
          final_master_log["Quarter"] = ""
          final_master_log["System_Flag"] = "Active"

        if not final_master_log.empty:
          if "match_key" in final_master_log.columns:
            final_master_log = final_master_log.drop(columns=["match_key"])

          base_cols = [
              "System_Flag",
              "Quarter",
              "Error_Details",
              "Partner_Response",
              "Case_Status",
          ]
          cols_order = base_cols + [
              c for c in final_master_log.columns if c not in base_cols
          ]
          final_master_log = final_master_log[cols_order]

          filter_view = st.radio(
              "Filter Error View:",
              [
                  "All Historical (Master Log)",
                  "Active Issues Only",
                  "Fixed/Resolved Only",
              ],
          )

          display_errors = final_master_log.copy()
          if filter_view == "Active Issues Only":
            display_errors = display_errors[
                display_errors["System_Flag"] == "Active"
            ]
            st.error(f"❌ {name}: {len(display_errors)} Active issues found.")
          elif filter_view == "Fixed/Resolved Only":
            display_errors = display_errors[
                display_errors["System_Flag"] == "Fixed/Resolved"
            ]
            st.success(
                f"✅ {name}: {len(display_errors)} issues have been resolved."
            )
          else:
            st.info(f"📁 {name}: {len(display_errors)} Total Historical records.")

          display_cols = [
              c for c in base_cols + ["CHILD_ID"] if c in display_errors.columns
          ]
          st.dataframe(
              display_errors[display_cols],
              use_container_width=True,
              hide_index=True,
          )

          st.download_button(
              f"📂 Download {name} Master Log",
              export_excel(final_master_log, f"{name}_Master_Log"),
              f"{name}_Master_Error_Log.xlsx",
              key=f"err_dl_{name}",
          )
      elif partner_rule_file:
        st.success(f"✅ {name} passed all rulebook checks!")
      else:
        st.info(f"Please upload {name} Rulebook to run checks.")
    else:
      st.info(f"No processed {name} data available.")

elif page == "EPIC Indicator":
  st.title("📊 EPIC Service Delivery Indicators")
  st.sidebar.header("📅 Reporting Filter")
  filter_type = st.sidebar.radio(
      "Select Filter Level", ["Monthly", "Quarterly"], key="epic_f_type"
  )

  current_year = datetime.now().year
  target_year = st.sidebar.number_input(
      "Select Year", value=current_year, key="epic_year_input"
  )

  if filter_type == "Quarterly":
    q_sel = st.sidebar.selectbox(
        "Select Quarter",
        ["Q1 (Jan-Mar)", "Q2 (Apr-Jun)", "Q3 (Jul-Sep)", "Q4 (Oct-Dec)"],
        key="epic_q_sel",
    )
    q_map = {
        "Q1 (Jan-Mar)": ["jan", "feb", "mar"],
        "Q2 (Apr-Jun)": ["apr", "may", "jun"],
        "Q3 (Jul-Sep)": ["jul", "aug", "sep"],
        "Q4 (Oct-Dec)": ["oct", "nov", "dec"],
    }
    target_months = q_map[q_sel]
    display_label = f"{q_sel} {target_year}"
  else:
    m_sel = st.sidebar.selectbox(
        "Select Month",
        [
            "January",
            "February",
            "March",
            "April",
            "May",
            "June",
            "July",
            "August",
            "September",
            "October",
            "November",
            "December",
        ],
        key="epic_m_sel",
    )
    target_months = [m_sel[:3].lower()]
    display_label = f"{m_sel} {target_year}"

  def calculate_indicators_epic(df, month_list):
    def parse_epic_age(age_val):
      if (
          pd.isna(age_val)
          or str(age_val).strip() == ""
          or str(age_val).lower() == "nan"
      ):
        return None, None
      s_val = str(age_val).strip()
      try:
        f_val = float(s_val)
        m = int(f_val)
        d = round((f_val - m) * 100)
        return m, d
      except ValueError:
        pass
      match = re.match(r"(\d+)m\s*(\d+)d", s_val, re.IGNORECASE)
      if match:
        return int(match.group(1)), int(match.group(2))
      return None, None

    def count_vax_gender(age_col, rp_col, min_age, max_age):
      if age_col not in df.columns or rp_col not in df.columns:
        return 0, 0, 0
      month_pattern = "|".join(month_list)
      vax_mask = (
          df[rp_col]
          .astype(str)
          .str.lower()
          .str.contains(month_pattern, na=False)
      )

      def row_filter(age_val):
        m, d = parse_epic_age(age_val)
        if m is None:
          return False
        decimal_val = m + (d / 100.0)
        return (decimal_val >= min_age) & (decimal_val < max_age)

      age_mask = df[age_col].apply(row_filter)
      filtered = df[vax_mask & age_mask]
      m = len(
          filtered[
              filtered["GENDER"]
              .astype(str)
              .str.lower()
              .str.startswith("m", na=False)
          ]
      )
      f = len(
          filtered[
              filtered["GENDER"]
              .astype(str)
              .str.lower()
              .str.startswith("f", na=False)
          ]
      )
      return m, f, len(filtered)

    ind_list = [
        (
            "Output 3.1: Under 1 child (MCV1)",
            "MMR1_Age",
            "MMR1_REPORTING_PERIOD",
            9.0,
            12.0,
        ),
        (
            "Output 3.2: Under 1 child (Penta1)",
            "Penta1_Age",
            "Penta1_REPORTING_PERIOD",
            1.12,
            12.0,
        ),
        (
            "Output 3.3: Under 1 child (Penta3)",
            "Penta3_Age",
            "Penta3_REPORTING_PERIOD",
            3.12,
            12.0,
        ),
        (
            "Output 3.4: Under 5 children (MCV1)",
            "MMR1_Age",
            "MMR1_REPORTING_PERIOD",
            9.0,
            60.0,
        ),
        (
            "Output 3.5: Under 5 children (Penta1)",
            "Penta1_Age",
            "Penta1_REPORTING_PERIOD",
            1.12,
            60.0,
        ),
        (
            "Output 3.6: Under 5 children (Penta3)",
            "Penta3_Age",
            "Penta3_REPORTING_PERIOD",
            3.12,
            60.0,
        ),
    ]
    rows = []
    for label, age_c, rp_c, amin, amax in ind_list:
      m, f, t = count_vax_gender(age_c, rp_c, amin, amax)
      rows.append({
          "Indicator": label,
          "Male": int(m),
          "Female": int(f),
          "Total": int(t),
      })
    return pd.DataFrame(rows)

  st.header(f"🏢 EPIC Data for {display_label}")
  partners = [("CMHN", "cmhn_processed")]
  for p_name, p_key in partners:
    full_key = f"{p_key}_v_final"
    if full_key not in st.session_state and p_key in st.session_state:
      final_v, _ = process_data_validation(st.session_state[p_key])
      st.session_state[full_key] = final_v

    if full_key in st.session_state and st.session_state[full_key] is not None:
      p_data = st.session_state[full_key].copy()
      y_mask = (
          pd.to_numeric(p_data["REPORTING_YEAR"], errors="coerce")
          == target_year
      )
      f_df = p_data[y_mask]
      if not f_df.empty:
        res = calculate_indicators_epic(f_df, target_months)
        st.dataframe(res, use_container_width=True, hide_index=True)
      else:
        st.warning(f"No records found for {target_year}.")